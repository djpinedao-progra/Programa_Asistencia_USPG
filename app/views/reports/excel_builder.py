"""Generación de libros Excel a partir de datos preparados."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font


def build_expo_workbook(category, categories, counts, attendees):
    """Construye el libro Excel del evento con conteos y asistentes ya consultados."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    summary.append(["Expo San Pablo", "Asistencia registrada"])
    summary.append(["Total general", sum(counts.values())])
    for key, label in categories.items():
        summary.append([label, counts.get(key, 0)])
    summary.append(["Filtro del detalle", categories.get(category, "Todos")])
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 28

    sheet = workbook.create_sheet("Asistentes y opiniones")
    sheet.append(
        [
            "Nombre",
            "Correo",
            "Categoría",
            "Participación",
            "Opinión",
            "Fecha y hora (Guatemala)",
        ]
    )
    for attendee in attendees:
        sheet.append(
            [
                attendee["name"],
                attendee["email"],
                categories[attendee["category"]],
                attendee["participation"],
                attendee["opinion"],
                attendee["recorded_at"],
            ]
        )
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in zip("ABCDEF", [30, 38, 20, 55, 70, 28]):
        sheet.column_dimensions[index].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
